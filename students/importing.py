"""Shared Excel-import logic for the Transaction table.

Used by both the `import_transactions` management command (a fixed path
on disk) and the in-app upload view (a file the user just posted) so the
two never drift apart.
"""
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.utils import timezone

from .models import Demand, Transaction

CURRENT_FEE_COLUMNS = [
    'SDF FEE', 'SDC FEE', 'AU FEE', 'Exam Fee', 'Other Fee',
    'Installment fee', 'Late Fee', 'Re-Admission Fee', 'Tuition Fee',
]
# Not a simple f'PAST DUES - {name}' derivation: the source spreadsheet
# itself is inconsistent — every other column matches its current-fee
# name exactly, but this one is "PAST DUES - Late fee" (lowercase "fee")
# where the current-fee column is "Late Fee". Hardcoded to match exactly
# what's actually in transactions.xlsx.
PAST_DUE_COLUMNS = [
    'PAST DUES - SDF FEE', 'PAST DUES - SDC FEE', 'PAST DUES - AU FEE', 'PAST DUES - Exam Fee',
    'PAST DUES - Other Fee', 'PAST DUES - Installment fee', 'PAST DUES - Late fee',
    'PAST DUES - Re-Admission Fee', 'PAST DUES - Tuition Fee',
]

REQUIRED_COLUMNS = ['Order', 'Name', 'Roll No']

HEADER_PREFIX = [
    'Transaction Date', 'Order', 'Tracking', 'Order Status', 'Status Message', 'App Code',
    'Amount', 'Merchant Amount', 'Currency', 'Billing Tel', 'Billing Email', 'Bank Name',
    'ID No.', 'Mobile Number', 'Email', 'Parent Mobile No', 'Admission Type', 'Fee Type',
    'Name', 'Group', 'Course', 'Course Fee', 'Section', 'Year', 'Roll No', 'Academic Year',
    'Issue Date', 'Due Date Installment 1',
]
HEADER_SUFFIX = ['Total', 'remarks']

# The full 48-column header row, built from the same pieces import_workbook()
# reads — so a downloaded sample file can never drift out of sync with what
# the importer actually expects.
ALL_HEADERS = HEADER_PREFIX + CURRENT_FEE_COLUMNS + PAST_DUE_COLUMNS + HEADER_SUFFIX

SAMPLE_ROW = {
    'Transaction Date': '2026-08-05 10:15:17',
    'Order': '7991741',
    'Tracking': '114712300800',
    'Order Status': 'Success',
    'Status Message': 'Transaction Successful-NA-0',
    'App Code': '713532414943',
    'Amount': 13525,
    'Merchant Amount': 13525,
    'Currency': 'INR',
    'Billing Tel': '9999999999',
    'Billing Email': 'student@example.com',
    'Bank Name': 'Union Bank of India',
    'ID No.': '76294020203',
    'Mobile Number': '9999999999',
    'Email': 'student@example.com',
    'Parent Mobile No': '8888888888',
    'Admission Type': 'CQ-LE',
    'Fee Type': 'TUTION FEE',
    'Name': 'SAMPLE STUDENT NAME',
    'Group': 'ECE',
    'Course': 'B.TECH',
    'Course Fee': 60225,
    'Section': '1',
    'Year': '2',
    'Roll No': '76294020203',
    'Academic Year': '2026-2027',
    'Issue Date': '01-08-2026',
    'Due Date Installment 1': '05-08-2026',
    'SDF FEE': 0, 'SDC FEE': 4000, 'AU FEE': 0, 'Exam Fee': 0, 'Other Fee': 0,
    'Installment fee': 0, 'Late Fee': 0, 'Re-Admission Fee': 0, 'Tuition Fee': 9525,
    'PAST DUES - SDF FEE': 0, 'PAST DUES - SDC FEE': 0, 'PAST DUES - AU FEE': 0,
    'PAST DUES - Exam Fee': 0, 'PAST DUES - Other Fee': 0, 'PAST DUES - Installment fee': 0,
    'PAST DUES - Late fee': 0, 'PAST DUES - Re-Admission Fee': 0, 'PAST DUES - Tuition Fee': 0,
    'Total': 13525,
    'remarks': '2ND YEAR FEE PAY BEFORE DUE DATE (LATERAL ENTRY)',
}


def build_sample_workbook():
    """An .xlsx with the exact header row import_workbook() expects, plus
    one filled-in example row, for staff to copy and fill in."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'sheet1'
    ws.append(ALL_HEADERS)
    ws.append([SAMPLE_ROW.get(col, '') for col in ALL_HEADERS])
    for col_idx, header in enumerate(ALL_HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(header) + 2)
    return wb


class ImportError_(Exception):
    """Raised for anything wrong with the workbook itself (not a single bad row)."""


def to_decimal(value):
    if value in (None, ''):
        return Decimal('0')
    try:
        return Decimal(str(value).replace(',', '').strip())
    except InvalidOperation:
        return Decimal('0')


def to_datetime(value):
    if not value:
        return None
    dt = value if isinstance(value, datetime) else None
    if dt is None:
        for fmt in ('%Y-%m-%d %H:%M:%S', '%d-%m-%Y %H:%M:%S'):
            try:
                dt = datetime.strptime(str(value).strip(), fmt)
                break
            except ValueError:
                continue
    if dt is None:
        return None
    return timezone.make_aware(dt) if timezone.is_naive(dt) else dt


def to_date(value):
    dt = to_datetime(value)
    if dt:
        return dt.date()
    if not value:
        return None
    for fmt in ('%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def import_workbook(wb, sheet_name=None):
    """Import every row of the given openpyxl Workbook into Transaction.

    Returns {'created': int, 'updated': int, 'skipped': int, 'total': int}.
    Raises ImportError_ if the sheet is empty or missing expected columns —
    that's a whole-file problem, distinct from a single skipped row.
    """
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ImportError_('That sheet is empty.')

    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ImportError_(
            f"This doesn't look like the transactions format — missing column(s): {', '.join(missing)}."
        )

    created, updated, skipped = 0, 0, 0

    for raw_row in rows[1:]:
        record = dict(zip(headers, raw_row))
        order_id = str(record.get('Order') or '').strip()
        if not order_id:
            skipped += 1
            continue

        current_fees = {col: str(record.get(col) or '0') for col in CURRENT_FEE_COLUMNS}
        past_dues = {col: str(record.get(col) or '0') for col in PAST_DUE_COLUMNS}

        defaults = dict(
            tracking=str(record.get('Tracking') or ''),
            transaction_date=to_datetime(record.get('Transaction Date')),
            order_status=str(record.get('Order Status') or ''),
            status_message=str(record.get('Status Message') or ''),
            app_code=str(record.get('App Code') or ''),
            amount=to_decimal(record.get('Amount')),
            merchant_amount=to_decimal(record.get('Merchant Amount')),
            currency=str(record.get('Currency') or ''),
            billing_tel=str(record.get('Billing Tel') or ''),
            billing_email=str(record.get('Billing Email') or ''),
            bank_name=str(record.get('Bank Name') or ''),
            id_no=str(record.get('ID No.') or ''),
            mobile_number=str(record.get('Mobile Number') or ''),
            email=str(record.get('Email') or ''),
            parent_mobile=str(record.get('Parent Mobile No') or ''),
            admission_type=str(record.get('Admission Type') or ''),
            fee_type=str(record.get('Fee Type') or ''),
            name=str(record.get('Name') or '').strip(),
            branch=str(record.get('Group') or '').strip(),
            course=str(record.get('Course') or '').strip(),
            course_fee=to_decimal(record.get('Course Fee')) if record.get('Course Fee') not in (None, '') else None,
            section=str(record.get('Section') or '').strip(),
            year=str(record.get('Year') or '').strip(),
            roll_no=str(record.get('Roll No') or '').strip(),
            academic_year=str(record.get('Academic Year') or '').strip(),
            issue_date=to_date(record.get('Issue Date')),
            due_date_installment1=to_date(record.get('Due Date Installment 1')),
            tuition_fee=to_decimal(record.get('Tuition Fee')),
            total=to_decimal(record.get('Total')),
            remarks=str(record.get('remarks') or ''),
            current_fees=current_fees,
            past_dues=past_dues,
        )

        _, was_created = Transaction.objects.update_or_create(order_id=order_id, defaults=defaults)
        created += was_created
        updated += not was_created

    return {'created': created, 'updated': updated, 'skipped': skipped, 'total': created + updated}


# --- Demand (upload.xlsx) -----------------------------------------------
#
# A different source format from transactions.xlsx: no Order/payment
# columns at all, and every sheet is a separate installment batch (e.g.
# "2ND YEAR 1st Install", "3RD YEAR 1st Intall") that all get imported —
# see import_demand_workbook(). It also has its own header quirk: unlike
# transactions.xlsx, THIS file spells the current-fee column "Late fee"
# (lowercase) and the past-dues one the same way, so a fresh column list
# is used rather than reusing CURRENT_FEE_COLUMNS above.

DEMAND_CURRENT_FEE_COLUMNS = [
    'SDF FEE', 'SDC FEE', 'AU FEE', 'Exam Fee', 'Other Fee',
    'Installment fee', 'Late fee', 'Re-Admission Fee', 'Tuition Fee',
]
DEMAND_PAST_DUE_COLUMNS = [f'PAST DUES - {name}' for name in DEMAND_CURRENT_FEE_COLUMNS]

DEMAND_REQUIRED_COLUMNS = ['Roll No', 'Name']

DEMAND_HEADER_PREFIX = [
    'ID No.', 'Mobile Number', 'Email', 'Parent Mobile No', 'Admission Type', 'Fee Type',
    'Name', 'Group', 'Course', 'Course Fee', 'Section', 'Year', 'Roll No', 'Academic Year',
    'Issue Date', 'Due Date Installment 1',
]
DEMAND_HEADER_SUFFIX = ['Total', 'Remarks']
DEMAND_ALL_HEADERS = DEMAND_HEADER_PREFIX + DEMAND_CURRENT_FEE_COLUMNS + DEMAND_PAST_DUE_COLUMNS + DEMAND_HEADER_SUFFIX

# One sample sheet per real installment batch — 2ND YEAR has no RTF
# column, the other two do — so the sample demonstrates both shapes
# instead of only the simplest one.
DEMAND_SAMPLE_SHEETS = [
    ('2ND YEAR 1st Install', False),
    ('3RD YEAR 1st Intall', True),
    ('4th Year 1st Install', True),
]

DEMAND_SAMPLE_ROWS = {
    '2ND YEAR 1st Install': {
        'ID No.': '325136408001', 'Mobile Number': '9490506907', 'Email': 'student@example.com',
        'Parent Mobile No': '9490506907', 'Admission Type': 'CQ', 'Fee Type': 'TUTION FEE',
        'Name': 'SAMPLE STUDENT NAME', 'Group': 'CIVIL', 'Course': 'B.TECH', 'Course Fee': 60225,
        'Section': 'A', 'Year': '2', 'Roll No': '325136408001', 'Academic Year': '2026-2027',
        'Issue Date': '01-07-2026', 'Due Date Installment 1': '20-07-2026',
        'SDF FEE': 0, 'SDC FEE': 4000, 'AU FEE': 0, 'Exam Fee': 0, 'Other Fee': 0,
        'Installment fee': 0, 'Late fee': 0, 'Re-Admission Fee': 0, 'Tuition Fee': 28212.5,
        'PAST DUES - SDF FEE': 0, 'PAST DUES - SDC FEE': 0, 'PAST DUES - AU FEE': 0,
        'PAST DUES - Exam Fee': 0, 'PAST DUES - Other Fee': 0, 'PAST DUES - Installment fee': 0,
        'PAST DUES - Late fee': 0, 'PAST DUES - Re-Admission Fee': 0, 'PAST DUES - Tuition Fee': 0,
        'Total': 32212.5, 'Remarks': '2ND YEAR FEE PAY BEFORE DUE DATE',
    },
    '3RD YEAR 1st Intall': {
        'ID No.': '324136408001', 'Mobile Number': '9391576007', 'Email': 'student@example.com',
        'Parent Mobile No': '9177901275', 'Admission Type': 'CQ SPOT', 'Fee Type': 'TUTION FEE',
        'Name': 'SAMPLE STUDENT NAME', 'Group': 'CIVIL', 'Course': 'B.TECH', 'Course Fee': 60025,
        'Section': 'A', 'Year': '3', 'Roll No': '324136408001', 'Academic Year': '2026-2027',
        'Issue Date': '01-07-2026', 'Due Date Installment 1': '20-07-2026',
        'SDF FEE': 0, 'SDC FEE': 3800, 'AU FEE': 0, 'Exam Fee': 0, 'Other Fee': 0,
        'Installment fee': 0, 'Late fee': 0, 'Re-Admission Fee': 0, 'Tuition Fee': 28212.5,
        'PAST DUES - SDF FEE': 0, 'PAST DUES - SDC FEE': 0, 'PAST DUES - AU FEE': 0,
        'PAST DUES - Exam Fee': 0, 'PAST DUES - Other Fee': 0, 'PAST DUES - Installment fee': 0,
        'PAST DUES - Late fee': 0, 'PAST DUES - Re-Admission Fee': 0, 'PAST DUES - Tuition Fee': 0,
        'Total': 32012.5, 'Remarks': 'Pay the 3rd Year fee before the due date', 'RTF': '',
    },
    '4th Year 1st Install': {
        'ID No.': '323136408001', 'Mobile Number': '9573346614', 'Email': 'student@example.com',
        'Parent Mobile No': '9347021056', 'Admission Type': 'CQ', 'Fee Type': 'TUTION FEE',
        'Name': 'SAMPLE STUDENT NAME', 'Group': 'CIVIL', 'Course': 'B.TECH', 'Course Fee': 55725,
        'Section': 'A', 'Year': '4', 'Roll No': '323136408001', 'Academic Year': '2026-2027',
        'Issue Date': '01-07-2026', 'Due Date Installment 1': '20-07-2026',
        'SDF FEE': 0, 'SDC FEE': 3200, 'AU FEE': 0, 'Exam Fee': 0, 'Other Fee': 0,
        'Installment fee': 0, 'Late fee': 0, 'Re-Admission Fee': 0, 'Tuition Fee': 26362.5,
        'PAST DUES - SDF FEE': 0, 'PAST DUES - SDC FEE': 0, 'PAST DUES - AU FEE': 0,
        'PAST DUES - Exam Fee': 0, 'PAST DUES - Other Fee': 0, 'PAST DUES - Installment fee': 0,
        'PAST DUES - Late fee': 0, 'PAST DUES - Re-Admission Fee': 0, 'PAST DUES - Tuition Fee': 0,
        'Total': 29562.5, 'Remarks': 'Pay the 4th year fee before the due date along with Past Dues', 'RTF': '',
    },
}


def build_demand_sample_workbook():
    """An .xlsx shaped like upload.xlsx: one sheet per installment batch,
    each with the exact header row import_demand_sheet() expects plus
    one filled-in example row, for staff to copy and fill in."""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, has_rtf in DEMAND_SAMPLE_SHEETS:
        headers = DEMAND_ALL_HEADERS + (['RTF'] if has_rtf else [])
        ws = wb.create_sheet(title=sheet_name)
        ws.append(headers)
        sample = DEMAND_SAMPLE_ROWS[sheet_name]
        ws.append([sample.get(col, '') for col in headers])
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(header) + 2)
    return wb


def to_str_id(value):
    """Roll No / ID No. in upload.xlsx are stored as numbers (e.g.
    325136408001.0) rather than text — str() on that gives '...001.0'.
    Strip the trailing '.0' for any whole-number float before stringifying."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_demand_header_row(rows):
    """upload.xlsx sheets aren't all aligned the same way — one has a
    blank row above the header. Scan for the row that actually looks
    like the header instead of assuming row 1."""
    for i, row in enumerate(rows):
        cells = [str(c).strip() if c is not None else '' for c in row]
        if 'Roll No' in cells and 'Name' in cells:
            return i, cells
    return None, None


def import_demand_sheet(ws):
    """Import one sheet of upload.xlsx into Demand, tagged with that
    sheet's name as the installment. Returns the same shape as
    import_workbook()'s result."""
    rows = list(ws.iter_rows(values_only=True))
    header_idx, headers = _find_demand_header_row(rows)
    if headers is None:
        raise ImportError_(
            f"Sheet '{ws.title}': couldn't find a header row with Roll No / Name columns."
        )
    missing = [col for col in DEMAND_REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ImportError_(
            f"Sheet '{ws.title}': missing column(s): {', '.join(missing)}."
        )

    created, updated, skipped = 0, 0, 0

    for raw_row in rows[header_idx + 1:]:
        record = dict(zip(headers, raw_row))
        roll_no = to_str_id(record.get('Roll No'))
        name = str(record.get('Name') or '').strip()
        if not roll_no or not name:
            skipped += 1
            continue

        current_fees = {col: str(record.get(col) or '0') for col in DEMAND_CURRENT_FEE_COLUMNS}
        past_dues = {col: str(record.get(col) or '0') for col in DEMAND_PAST_DUE_COLUMNS}

        defaults = dict(
            id_no=to_str_id(record.get('ID No.')),
            mobile_number=to_str_id(record.get('Mobile Number')),
            email=str(record.get('Email') or '').strip(),
            parent_mobile=to_str_id(record.get('Parent Mobile No')),
            admission_type=str(record.get('Admission Type') or '').strip(),
            fee_type=str(record.get('Fee Type') or '').strip(),
            name=name,
            branch=str(record.get('Group') or '').strip(),
            course=str(record.get('Course') or '').strip(),
            course_fee=to_decimal(record.get('Course Fee')) if record.get('Course Fee') not in (None, '') else None,
            section=str(record.get('Section') or '').strip(),
            year=to_str_id(record.get('Year')),
            academic_year=str(record.get('Academic Year') or '').strip(),
            issue_date=to_date(record.get('Issue Date')),
            due_date_installment1=to_date(record.get('Due Date Installment 1')),
            total=to_decimal(record.get('Total')),
            remarks=str(record.get('Remarks') or '').strip(),
            rtf=to_decimal(record.get('RTF')) if record.get('RTF') not in (None, '') else None,
            current_fees=current_fees,
            past_dues=past_dues,
        )

        _, was_created = Demand.objects.update_or_create(
            roll_no=roll_no, installment=ws.title, defaults=defaults,
        )
        created += was_created
        updated += not was_created

    return {'created': created, 'updated': updated, 'skipped': skipped, 'total': created + updated}


def import_demand_workbook(wb):
    """Import every sheet of upload.xlsx into Demand — each sheet is a
    separate installment batch and all of them get imported, not just
    the first. Returns the combined totals plus a per-sheet breakdown."""
    combined = {'created': 0, 'updated': 0, 'skipped': 0, 'total': 0, 'sheets': []}
    for sheet_name in wb.sheetnames:
        result = import_demand_sheet(wb[sheet_name])
        combined['created'] += result['created']
        combined['updated'] += result['updated']
        combined['skipped'] += result['skipped']
        combined['total'] += result['total']
        combined['sheets'].append({'name': sheet_name, **result})
    return combined
