import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Max, Q, Sum
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render

from .decorators import office_required
from .departments import DEPARTMENT_CODE_CHOICES, DEPARTMENTS_BY_SLUG, STATUS_CHOICES, user_can_act_for
from .forms import (
    ApprovalDecisionForm, DashboardFilterForm, DemandFilterForm, DemandUploadForm, FeeRowFormSet,
    NoDuesForm, StudentFilterForm, TCApplicationForm, TransactionUploadForm,
)
from .importing import (
    ImportError_, build_demand_sample_workbook, build_sample_workbook,
    import_demand_workbook, import_workbook,
)
from .models import Clearance, Demand, DueApproval, FeePaidRow, Transaction


def _annotate_demand_payment_status(demands):
    """Mutates each Demand in place, adding .paid_total / .is_paid — a
    demand row counts as paid if that roll_no has a transaction for the
    same year of study. Shared by the dashboard and the Student Demand
    list so the two never disagree on what "paid" means."""
    roll_nos = [d.roll_no for d in demands]
    paid_totals = {}
    for row in (
        Transaction.objects.filter(roll_no__in=roll_nos)
        .values('roll_no', 'year').annotate(paid_total=Sum('total'))
    ):
        paid_totals[(row['roll_no'], row['year'])] = row['paid_total']
    for d in demands:
        d.paid_total = paid_totals.get((d.roll_no, d.year))
        d.is_paid = d.paid_total is not None


def _rate_accent(rate):
    if rate >= 75:
        return 'success'
    if rate >= 40:
        return 'warning'
    return 'danger'


def _filtered_querysets(request):
    """Applies the dashboard's branch/course/academic_year/installment
    filters (from the top filter row, or from a modal's own request) to
    fresh Transaction/Demand querysets. Shared by the dashboard itself,
    its exports, and the branch-detail modal/export — one definition of
    "the current filter" everything else builds on."""
    form = DashboardFilterForm(request.GET or None)

    txn_qs = Transaction.objects.all()
    demand_qs = Demand.objects.all()

    if form.is_valid():
        data = form.cleaned_data
        if data.get('branch'):
            txn_qs = txn_qs.filter(branch=data['branch'])
            demand_qs = demand_qs.filter(branch=data['branch'])
        if data.get('course'):
            txn_qs = txn_qs.filter(course=data['course'])
            demand_qs = demand_qs.filter(course=data['course'])
        if data.get('academic_year'):
            txn_qs = txn_qs.filter(academic_year=data['academic_year'])
            demand_qs = demand_qs.filter(academic_year=data['academic_year'])
        if data.get('installment'):
            demand_qs = demand_qs.filter(installment=data['installment'])

    return form, txn_qs, demand_qs


def _dashboard_context(request):
    """Everything the dashboard page (and its Excel exports) need, built
    from the same filtered querysets — so an export always matches
    exactly what's on screen, never a different slice of the data."""
    form, txn_qs, demand_qs = _filtered_querysets(request)

    total_collected = txn_qs.aggregate(v=Sum('total'))['v'] or 0
    total_students = txn_qs.values('roll_no').distinct().count()

    demands = list(demand_qs)
    _annotate_demand_payment_status(demands)

    total_demand = sum(float(d.total or 0) for d in demands)
    paid_demand_amount = sum(float(d.total or 0) for d in demands if d.is_paid)
    pending_amount = total_demand - paid_demand_amount
    paid_students = sum(1 for d in demands if d.is_paid)
    unpaid_students = len(demands) - paid_students
    collection_rate = round(paid_demand_amount / total_demand * 100, 1) if total_demand else 0

    def _grouped(key_func):
        groups = {}
        for d in demands:
            key = key_func(d) or 'Unspecified'
            row = groups.setdefault(key, {'demand': 0.0, 'paid': 0.0, 'pending': 0.0, 'students': 0, 'paid_students': 0})
            amount = float(d.total or 0)
            row['demand'] += amount
            row['students'] += 1
            if d.is_paid:
                row['paid'] += amount
                row['paid_students'] += 1
            else:
                row['pending'] += amount
        for row in groups.values():
            row['rate'] = round(row['paid'] / row['demand'] * 100, 1) if row['demand'] else 0
        return groups

    by_branch = dict(sorted(_grouped(lambda d: d.branch).items(), key=lambda kv: -kv[1]['demand']))
    # Installment names happen to sort correctly as plain strings
    # ("2ND YEAR…" < "3RD YEAR…" < "4th Year…") since they lead with the
    # year digit — no separate ordering key needed.
    by_installment = dict(sorted(_grouped(lambda d: d.installment).items()))

    # Full (untruncated) lists for the Excel exports; the dashboard page
    # itself only shows the first page of each as a preview.
    all_pending = sorted((d for d in demands if not d.is_paid), key=lambda d: -(d.total or 0))
    all_recent = list(txn_qs.order_by('-transaction_date'))
    student_paid_rate = round(paid_students / len(demands) * 100, 1) if demands else 0

    return {
        'form': form,
        'total_students': total_students,
        'total_collected': total_collected,
        'total_demand': total_demand,
        'paid_demand_amount': paid_demand_amount,
        'pending_amount': pending_amount,
        'collection_rate': collection_rate,
        'collection_rate_accent': _rate_accent(collection_rate),
        'total_demand_students': len(demands),
        'paid_students': paid_students,
        'unpaid_students': unpaid_students,
        'student_paid_rate': student_paid_rate,
        'student_paid_rate_accent': _rate_accent(student_paid_rate),
        'by_branch': by_branch,
        'by_installment': by_installment,
        'top_pending': all_pending[:8],
        'recent': all_recent[:8],
        'all_pending': all_pending,
        'all_recent': all_recent,
        'branch_chart_data': {
            'labels': list(by_branch.keys()),
            'paid': [round(row['paid'], 2) for row in by_branch.values()],
            'pending': [round(row['pending'], 2) for row in by_branch.values()],
        },
        'installment_chart_data': {
            'labels': list(by_installment.keys()),
            'paid': [round(row['paid'], 2) for row in by_installment.values()],
            'pending': [round(row['pending'], 2) for row in by_installment.values()],
        },
    }


@login_required
@office_required
def dashboard(request):
    return render(request, 'dashboard.html', _dashboard_context(request))


def _xlsx_response(filename, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(str(header)) + 2)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@office_required
def dashboard_export_branch(request):
    ctx = _dashboard_context(request)
    headers = ['Branch', 'Students', 'Paid Students', 'Collected (Rs.)', 'Demand (Rs.)', 'Pending (Rs.)', 'Rate (%)']
    rows = [
        [branch, row['students'], row['paid_students'], row['paid'], row['demand'], row['pending'], row['rate']]
        for branch, row in ctx['by_branch'].items()
    ]
    return _xlsx_response('branch_wise_breakdown.xlsx', headers, rows)


@login_required
@office_required
def dashboard_export_pending(request):
    ctx = _dashboard_context(request)
    headers = ['Name', 'Roll No', 'Course', 'Branch', 'Section', 'Year', 'Academic Year', 'Installment', 'Pending (Rs.)']
    rows = [
        [d.name, d.roll_no, d.course, d.branch, d.section, d.year, d.academic_year, d.installment, float(d.total or 0)]
        for d in ctx['all_pending']
    ]
    return _xlsx_response('needs_attention_pending.xlsx', headers, rows)


@login_required
@office_required
def dashboard_export_recent(request):
    ctx = _dashboard_context(request)
    headers = ['Name', 'Roll No', 'Course', 'Branch', 'Amount (Rs.)', 'Date']
    rows = [
        [
            t.name, t.roll_no, t.course, t.branch, float(t.total or 0),
            t.transaction_date.strftime('%Y-%m-%d %H:%M') if t.transaction_date else '',
        ]
        for t in ctx['all_recent']
    ]
    return _xlsx_response('recent_transactions.xlsx', headers, rows)


def _branch_modal_data(request):
    """Shared by the branch-detail modal (JSON) and its Excel export:
    the current dashboard filters, narrowed to one branch and one
    payment-status slice — 'all' / 'paid' / 'unpaid' — as clicked from
    a Branch-wise Breakdown cell."""
    _, txn_qs, demand_qs = _filtered_querysets(request)
    status = request.GET.get('status', 'all')

    demands = list(demand_qs)
    _annotate_demand_payment_status(demands)
    if status == 'paid':
        demands = [d for d in demands if d.is_paid]
    elif status == 'unpaid':
        demands = [d for d in demands if not d.is_paid]
    demands.sort(key=lambda d: d.name)

    transactions = list(txn_qs.order_by('-transaction_date'))
    return status, demands, transactions


@login_required
@office_required
def dashboard_branch_detail(request):
    status, demands, transactions = _branch_modal_data(request)
    return JsonResponse({
        'branch': request.GET.get('branch') or 'All Branches',
        'status': status,
        'demand_rows': [
            {
                'name': d.name, 'roll_no': d.roll_no, 'course': d.course,
                'section': d.section, 'year': d.year, 'academic_year': d.academic_year,
                'installment': d.installment, 'total': float(d.total or 0), 'is_paid': d.is_paid,
            }
            for d in demands
        ],
        'transaction_rows': [
            {
                'name': t.name, 'roll_no': t.roll_no, 'order_id': t.order_id,
                'total': float(t.total or 0),
                'date': t.transaction_date.strftime('%d %b %Y') if t.transaction_date else '',
                'status': t.order_status,
            }
            for t in transactions
        ],
    })


@login_required
@office_required
def dashboard_branch_export(request):
    status, demands, transactions = _branch_modal_data(request)
    branch = request.GET.get('branch') or 'all'

    wb = openpyxl.Workbook()
    ws_demand = wb.active
    ws_demand.title = 'Demand'
    ws_demand.append(['Name', 'Roll No', 'Course', 'Section', 'Year', 'Academic Year', 'Installment', 'Amount (Rs.)', 'Paid'])
    for d in demands:
        ws_demand.append([
            d.name, d.roll_no, d.course, d.section, d.year, d.academic_year,
            d.installment, float(d.total or 0), 'Yes' if d.is_paid else 'No',
        ])

    ws_txn = wb.create_sheet('Transactions')
    ws_txn.append(['Name', 'Roll No', 'Order ID', 'Amount (Rs.)', 'Date', 'Status'])
    for t in transactions:
        ws_txn.append([
            t.name, t.roll_no, t.order_id, float(t.total or 0),
            t.transaction_date.strftime('%Y-%m-%d %H:%M') if t.transaction_date else '', t.order_status,
        ])

    for ws in (ws_demand, ws_txn):
        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(12, len(str(header)) + 2)

    filename = f'{branch}_{status}_detail.xlsx'.replace(' ', '_')
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
@office_required
def student_list(request):
    form = StudentFilterForm(request.GET or None)
    qs = Transaction.objects.all()

    if form.is_valid():
        data = form.cleaned_data
        if data.get('q'):
            q = data['q']
            qs = qs.filter(
                Q(name__icontains=q) | Q(roll_no__icontains=q) | Q(id_no__icontains=q)
                | Q(order_id__icontains=q) | Q(app_code__icontains=q)
            )
        for field in ('course', 'branch', 'section', 'year', 'academic_year', 'admission_type'):
            if data.get(field):
                qs = qs.filter(**{field: data[field]})

    students = list(
        qs.values('roll_no')
        .annotate(
            name=Max('name'), course=Max('course'), branch=Max('branch'),
            section=Max('section'), year=Max('year'), academic_year=Max('academic_year'),
            mobile_number=Max('mobile_number'), total_paid=Sum('total'),
            last_payment=Max('transaction_date'), payments=Count('id'),
        )
        .order_by('name')
    )

    # One query for every Clearance + its approvals on this page, rather
    # than one per row — a roll_no with no Clearance yet (workflow never
    # sent) just falls back to NOT_SENT.
    clearances = {
        c.roll_no: c
        for c in Clearance.objects.filter(roll_no__in=[s['roll_no'] for s in students]).prefetch_related('approvals')
    }
    for s in students:
        clearance = clearances.get(s['roll_no'])
        s['approval_status'] = clearance.approval_status() if clearance else 'NOT_SENT'

    return render(request, 'students/list.html', {'form': form, 'students': students})


@login_required
@office_required
def demand_list(request):
    """Every fee-demand row from upload.xlsx, cross-checked year-for-year
    against what's actually in Transaction — so office staff can see who
    still owes their installment and who's already paid it."""
    form = DemandFilterForm(request.GET or None)
    qs = Demand.objects.all()

    status_filter = ''
    if form.is_valid():
        data = form.cleaned_data
        status_filter = data.get('status') or ''
        if data.get('q'):
            q = data['q']
            qs = qs.filter(Q(name__icontains=q) | Q(roll_no__icontains=q) | Q(id_no__icontains=q))
        for field in ('installment', 'course', 'branch', 'section', 'year', 'academic_year'):
            if data.get(field):
                qs = qs.filter(**{field: data[field]})

    demands = list(qs.order_by('name'))
    _annotate_demand_payment_status(demands)

    # Counted before the status filter is applied, so all three stat
    # cards always show what clicking them would give you — a card
    # never reports "0" just because a different one is currently active.
    total_count = len(demands)
    paid_count = sum(1 for d in demands if d.is_paid)
    unpaid_count = total_count - paid_count

    if status_filter == 'PAID':
        demands = [d for d in demands if d.is_paid]
    elif status_filter == 'UNPAID':
        demands = [d for d in demands if not d.is_paid]

    return render(request, 'students/demand_list.html', {
        'form': form, 'demands': demands, 'status_filter': status_filter,
        'total_count': total_count, 'paid_count': paid_count, 'unpaid_count': unpaid_count,
    })


@login_required
@office_required
def demand_upload(request):
    """Same idea as student_upload, but for the Demand table and driven
    entirely by AJAX (see templates/students/demand_upload.html) — POST
    always gets a JSON response back, never a redirect, so the page can
    show a spinner and swap in the result without a full reload."""
    if request.method == 'POST':
        form = DemandUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            error = form.errors.get('excel_file', ['Please choose a file to upload.'])[0]
            return JsonResponse({'success': False, 'error': error}, status=400)

        excel_file = form.cleaned_data['excel_file']
        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            result = import_demand_workbook(wb)
        except ImportError_ as exc:
            return JsonResponse({'success': False, 'error': str(exc)}, status=400)
        except Exception:
            return JsonResponse({
                'success': False,
                'error': "Couldn't read that file — make sure it's a valid .xlsx workbook.",
            }, status=400)

        message = (
            f"Imported {result['total']} demand records across {len(result['sheets'])} sheet(s) "
            f"({result['created']} new, {result['updated']} updated, {result['skipped']} skipped)."
        )
        return JsonResponse({'success': True, 'message': message, 'result': result})

    form = DemandUploadForm()
    return render(request, 'students/demand_upload.html', {'form': form})


@login_required
@office_required
def demand_sample_format(request):
    wb = build_demand_sample_workbook()
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="demand_sample_format.xlsx"'
    wb.save(response)
    return response


@login_required
@office_required
def student_upload(request):
    if request.method == 'POST':
        form = TransactionUploadForm(request.POST, request.FILES)
        if form.is_valid():
            excel_file = form.cleaned_data['excel_file']
            try:
                wb = openpyxl.load_workbook(excel_file, data_only=True)
                result = import_workbook(wb)
            except ImportError_ as exc:
                form.add_error('excel_file', str(exc))
            except Exception:
                form.add_error('excel_file', "Couldn't read that file — make sure it's a valid .xlsx workbook.")
            else:
                messages.success(
                    request,
                    f"Imported {result['total']} transactions "
                    f"({result['created']} new, {result['updated']} updated, {result['skipped']} skipped).",
                )
                return redirect('student-list')
    else:
        form = TransactionUploadForm()

    return render(request, 'students/upload.html', {'form': form})


@login_required
@office_required
def sample_format(request):
    wb = build_sample_workbook()
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="transactions_sample_format.xlsx"'
    wb.save(response)
    return response


def _student_context(roll_no):
    """Builds the detail-page context for any roll_no that appears
    *either* in Transaction (has paid something) or Demand (owes an
    installment but hasn't paid yet) — a demand-only student still gets
    a real detail page, just without the payment history / certificate
    section that depends on an actual transaction existing."""
    transactions = Transaction.objects.filter(roll_no=roll_no).order_by('-transaction_date')
    demands = Demand.objects.filter(roll_no=roll_no).order_by('-academic_year', '-year')
    has_transactions = transactions.exists()
    if not has_transactions and not demands.exists():
        return None

    latest = transactions.first() or demands.first()
    totals = transactions.aggregate(total_paid=Sum('total'))
    clearance, _ = Clearance.objects.get_or_create(roll_no=roll_no)

    # Same rule as the Student Demand list: a demand row counts as paid
    # if this roll_no has any transaction for that same year of study.
    paid_years = set(transactions.values_list('year', flat=True))
    demands = list(demands)
    for d in demands:
        d.is_paid = d.year in paid_years

    return {
        'roll_no': roll_no,
        'latest': latest,
        'transactions': transactions,
        'has_transactions': has_transactions,
        'demands': demands,
        'total_demand': sum((d.total or 0) for d in demands),
        'total_paid': totals['total_paid'] or 0,
        'clearance': clearance,
    }


def _effective_fields(clearance, latest):
    """Certificate display values: the manual override on Clearance if
    present, otherwise the value derived from the transaction table."""
    course_branch = f'{latest.course} / {latest.branch}'.strip(' /') if latest else ''
    year_of_study = f'Year {latest.year} · {latest.academic_year}' if latest else ''
    return {
        'student_name': clearance.student_name or (latest.name if latest else ''),
        'course_branch_name': clearance.course_branch_name or course_branch,
        'roll_no_register': clearance.roll_no_register or clearance.roll_no,
        'year_of_study': clearance.year_of_study or year_of_study,
        'section_display': clearance.section_display or (latest.section if latest else ''),
    }


@login_required
@office_required
def student_detail(request, roll_no):
    context = _student_context(roll_no)
    if context is None:
        return render(request, 'students/not_found.html', {'roll_no': roll_no}, status=404)
    clearance = context['clearance']
    context['approvals'] = clearance.ordered_approvals()
    context['overall_status'] = clearance.approval_status()
    return render(request, 'students/detail.html', context)


def _seed_fee_rows(clearance):
    """Add a FeePaidRow for any transaction not already represented —
    safe to call on every certificate view/edit, so a fee transaction
    uploaded after the certificate was first opened (e.g. next year's
    fee) still shows up next time, without touching rows already
    edited or added by hand."""
    already_synced = set(
        clearance.fee_rows.exclude(source_order_id='').values_list('source_order_id', flat=True)
    )
    new_transactions = (
        Transaction.objects.filter(roll_no=clearance.roll_no)
        .exclude(order_id__in=already_synced)
        .order_by('transaction_date')
    )
    if not new_transactions:
        return

    next_order = clearance.fee_rows.aggregate(Max('order'))['order__max']
    next_order = 0 if next_order is None else next_order + 1
    FeePaidRow.objects.bulk_create([
        FeePaidRow(
            clearance=clearance, order=next_order + i,
            label=f'Year {t.year} · {t.academic_year}',
            amount_paid=t.total,
            date_paid=t.transaction_date.date() if t.transaction_date else None,
            receipt_no=t.order_id,
            source_order_id=t.order_id,
        )
        for i, t in enumerate(new_transactions)
    ])


@login_required
@office_required
def send_for_approval(request, roll_no):
    """Kick off (or top up) the parallel approval workflow: create a
    PENDING DueApproval for every department that doesn't already have
    one for this student. All three are created together in this one
    call, so Library, Laboratory, and Sports are all notified/queued
    at the same time — not one after another."""
    if request.method != 'POST':
        return redirect('student-detail', roll_no=roll_no)

    latest = Transaction.objects.filter(roll_no=roll_no).order_by('-transaction_date').first()
    if latest is None:
        return render(request, 'students/not_found.html', {'roll_no': roll_no}, status=404)

    clearance, _ = Clearance.objects.get_or_create(roll_no=roll_no)
    existing = set(clearance.approvals.values_list('department', flat=True))
    sent_to = []
    for code, label in DEPARTMENT_CODE_CHOICES:
        if code not in existing:
            DueApproval.objects.create(clearance=clearance, department=code)
            sent_to.append(label)

    if sent_to:
        messages.success(request, f"Sent for approval to {', '.join(sent_to)} — awaiting their decision.")
    else:
        messages.info(request, 'Already sent for approval to all departments — awaiting their decision.')

    return redirect(request.POST.get('next') or 'student-detail', roll_no=roll_no)


@login_required
@office_required
def certificate_view(request, roll_no):
    context = _student_context(roll_no)
    if context is None:
        return render(request, 'students/not_found.html', {'roll_no': roll_no}, status=404)
    clearance = context['clearance']
    _seed_fee_rows(clearance)
    context['eff'] = _effective_fields(clearance, context['latest'])
    context['fee_rows'] = clearance.fee_rows.all()
    context['fee_total'] = sum((r.amount_paid or 0) for r in context['fee_rows'])
    context['approvals'] = clearance.ordered_approvals()
    context['overall_status'] = clearance.approval_status()
    return render(request, 'students/certificate.html', context)


@login_required
@office_required
def certificate_edit(request, roll_no):
    latest = Transaction.objects.filter(roll_no=roll_no).order_by('-transaction_date').first()
    if latest is None:
        return render(request, 'students/not_found.html', {'roll_no': roll_no}, status=404)
    clearance, _ = Clearance.objects.get_or_create(roll_no=roll_no)
    _seed_fee_rows(clearance)
    defaults = _effective_fields(clearance, latest)

    if request.method == 'POST':
        tc_form = TCApplicationForm(request.POST, instance=clearance, defaults=defaults, prefix='tc')
        no_dues_form = NoDuesForm(request.POST, instance=clearance, defaults=defaults, prefix='nodues')
        formset = FeeRowFormSet(request.POST, instance=clearance, prefix='feerows')
        if tc_form.is_valid() and no_dues_form.is_valid() and formset.is_valid():
            tc_form.save()
            no_dues_form.save()
            formset.save()
            return redirect('certificate-view', roll_no=roll_no)
    else:
        tc_form = TCApplicationForm(instance=clearance, defaults=defaults, prefix='tc')
        no_dues_form = NoDuesForm(instance=clearance, defaults=defaults, prefix='nodues')
        formset = FeeRowFormSet(instance=clearance, prefix='feerows')

    return render(request, 'students/certificate_form.html', {
        'tc_form': tc_form, 'no_dues_form': no_dues_form, 'formset': formset,
        'roll_no': roll_no, 'latest': latest,
        'approvals': clearance.ordered_approvals(), 'overall_status': clearance.approval_status(),
    })


@login_required
def approval_queue(request, dept_slug):
    """A department admin's own queue — every student sent to their
    department, filterable by decision status. Only members of that
    department's Group (or superusers) may view it."""
    dept = DEPARTMENTS_BY_SLUG.get(dept_slug)
    if dept is None:
        return render(request, 'students/not_found.html', {'roll_no': dept_slug}, status=404)
    if not user_can_act_for(request.user, dept):
        return HttpResponseForbidden("You don't have access to this department's approval queue.")

    status_filter = request.GET.get('status', 'PENDING')
    qs = DueApproval.objects.filter(department=dept['code']).select_related('clearance').order_by('-requested_at')
    if status_filter in dict(STATUS_CHOICES):
        qs = qs.filter(status=status_filter)

    # Per-roll_no "latest transaction", done in Python rather than
    # QuerySet.distinct(*fields) — that variant is Postgres-only and this
    # project runs on SQLite.
    latest_by_roll = {}
    roll_nos = [a.clearance.roll_no for a in qs]
    for t in Transaction.objects.filter(roll_no__in=roll_nos).order_by('roll_no', '-transaction_date'):
        latest_by_roll.setdefault(t.roll_no, t)

    rows = [
        {'approval': approval, 'latest': latest_by_roll.get(approval.clearance.roll_no)}
        for approval in qs
    ]
    counts = {
        code: DueApproval.objects.filter(department=dept['code'], status=code).count()
        for code, _ in STATUS_CHOICES
    }

    return render(request, 'students/approval_queue.html', {
        'dept': dept, 'rows': rows, 'status_filter': status_filter, 'counts': counts,
    })


@login_required
def approval_decide(request, dept_slug, approval_id):
    """A department admin's Cleared / Dues decision for one student,
    independent of what the other two departments decide."""
    dept = DEPARTMENTS_BY_SLUG.get(dept_slug)
    if dept is None:
        return render(request, 'students/not_found.html', {'roll_no': dept_slug}, status=404)
    if not user_can_act_for(request.user, dept):
        return HttpResponseForbidden("You don't have access to this department's approval queue.")

    approval = get_object_or_404(DueApproval, pk=approval_id, department=dept['code'])

    if request.method == 'POST':
        form = ApprovalDecisionForm(request.POST)
        if form.is_valid():
            approval.mark(
                form.cleaned_data['status'], request.user,
                form.cleaned_data['remarks'], form.cleaned_data['amount_due'],
            )
            messages.success(
                request,
                f"{approval.clearance.roll_no} marked {approval.get_status_display()} for {dept['label']}.",
            )
        else:
            messages.error(request, 'Please choose a valid decision.')

    return redirect('approval-queue', dept_slug=dept_slug)
